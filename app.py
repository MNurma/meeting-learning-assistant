import os
import io
import re
import time
import shutil
import stat
import tempfile

import streamlit as st
import whisper
import google.generativeai as genai
import imageio_ffmpeg
from pydub import AudioSegment
from docx import Document
from fpdf import FPDF


# =========================
# SETUP FFMPEG
# =========================

_ffmpeg_source = imageio_ffmpeg.get_ffmpeg_exe()
_ffmpeg_dir = os.path.join(tempfile.gettempdir(), "ffmpeg_bin")
os.makedirs(_ffmpeg_dir, exist_ok=True)
_ffmpeg_target = os.path.join(_ffmpeg_dir, "ffmpeg")

if not os.path.exists(_ffmpeg_target):
    shutil.copy(_ffmpeg_source, _ffmpeg_target)
    st_mode = os.stat(_ffmpeg_target).st_mode
    os.chmod(_ffmpeg_target, st_mode | stat.S_IEXEC | stat.S_IXGRP | stat.S_IXOTH)

os.environ["PATH"] = _ffmpeg_dir + os.pathsep + os.environ["PATH"]
AudioSegment.converter = _ffmpeg_target
AudioSegment.ffmpeg = _ffmpeg_target


# =========================
# KONFIGURASI HALAMAN
# =========================

st.set_page_config(
    page_title="Meeting Learning Assistant",
    page_icon="🎙️",
    layout="wide"
)

st.title("🎙️ Meeting Learning Assistant")
st.write(
    "Upload rekaman meeting, lalu ubah menjadi transkrip "
    "dan bahan belajar."
)


# =========================
# CEK API KEY GEMINI
# =========================

gemini_key = st.secrets.get("GEMINI_API_KEY") or os.getenv("GEMINI_API_KEY")

if not gemini_key:
    st.error(
        "GEMINI_API_KEY belum ditemukan. "
        "Silakan set API key terlebih dahulu."
    )
    st.stop()

genai.configure(api_key=gemini_key)


# =========================
# LOAD MODEL WHISPER (sekali saja, di-cache)
# =========================

@st.cache_resource
def load_whisper_model():
    return whisper.load_model("base")


# =========================
# HELPER: PARSING MARKDOWN SEDERHANA
# =========================

def parse_inline_formatting(text):
    """Pecah teks jadi list (teks, is_bold, is_italic).
    Menangani **bold** dan *italic*.
    """
    parts = []
    pattern = re.compile(r"\*\*(.+?)\*\*|\*(.+?)\*")
    last_end = 0

    for match in pattern.finditer(text):
        if match.start() > last_end:
            parts.append((text[last_end:match.start()], False, False))

        if match.group(1) is not None:
            parts.append((match.group(1), True, False))
        else:
            parts.append((match.group(2), False, True))

        last_end = match.end()

    if last_end < len(text):
        parts.append((text[last_end:], False, False))
    if not parts:
        parts.append((text, False, False))

    return parts


def sanitize_for_pdf(text):
    """Ganti karakter unicode yang tidak didukung font Helvetica dengan versi ASCII."""
    replacements = {
        "—": "-",
        "–": "-",
        "‘": "'",
        "’": "'",
        "“": '"',
        "”": '"',
        "…": "...",
        "•": "-",
        "\u00a0": " ",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text


# =========================
# FUNGSI GENERATE FILE DOWNLOAD
# =========================

def generate_docx(transcript, summary):
    doc = Document()
    doc.add_heading("Meeting Learning Assistant", level=1)

    doc.add_heading("Transkrip", level=2)
    doc.add_paragraph(transcript)

    doc.add_heading("Bahan Belajar", level=2)

    for raw_line in summary.split("\n"):
        line = raw_line.strip()

        if not line or line == "---":
            continue

        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
            continue

        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
            continue

        is_bullet = False
        if line.startswith("* ") or line.startswith("- "):
            line = line[2:].strip()
            is_bullet = True

        paragraph = doc.add_paragraph(style="List Bullet" if is_bullet else None)

        for text_part, is_bold, is_italic in parse_inline_formatting(line):
            run = paragraph.add_run(text_part)
            run.bold = is_bold
            run.italic = is_italic

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def generate_pdf(transcript, summary):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)

    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "Meeting Learning Assistant", ln=True)
    pdf.ln(5)

    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 10, "Transkrip", ln=True)
    pdf.set_font("Helvetica", size=11)
    pdf.multi_cell(0, 7, sanitize_for_pdf(transcript))
    pdf.ln(5)

    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 10, "Bahan Belajar", ln=True)
    pdf.ln(2)

    for raw_line in summary.split("\n"):
        line = raw_line.strip()

        if not line or line == "---":
            continue

        if line.startswith("### "):
            pdf.ln(2)
            pdf.set_font("Helvetica", "B", 12)
            pdf.multi_cell(0, 7, sanitize_for_pdf(line[4:].strip()))
            pdf.set_font("Helvetica", size=11)
            continue

        if line.startswith("## "):
            pdf.ln(2)
            pdf.set_font("Helvetica", "B", 13)
            pdf.multi_cell(0, 7, sanitize_for_pdf(line[3:].strip()))
            pdf.set_font("Helvetica", size=11)
            continue

        prefix = ""
        if line.startswith("* ") or line.startswith("- "):
            line = line[2:].strip()
            prefix = "- "

        parts = parse_inline_formatting(line)

        pdf.set_x(pdf.l_margin)
        if prefix:
            pdf.write(7, prefix)

        for text_part, is_bold, is_italic in parts:
            if is_bold and is_italic:
                style = "BI"
            elif is_bold:
                style = "B"
            elif is_italic:
                style = "I"
            else:
                style = ""
            pdf.set_font("Helvetica", style, 11)
            pdf.write(7, sanitize_for_pdf(text_part))

        pdf.ln(9)

    raw = pdf.output(dest="S")
    if isinstance(raw, str):
        raw = raw.encode("latin-1", errors="replace")
    else:
        raw = bytes(raw)

    buffer = io.BytesIO(raw)
    return buffer


# =========================
# INISIALISASI SESSION STATE
# =========================

if "transcript" not in st.session_state:
    st.session_state.transcript = None

if "result_text" not in st.session_state:
    st.session_state.result_text = None


# =========================
# UPLOAD AUDIO
# =========================

uploaded_file = st.file_uploader(
    "Upload rekaman meeting",
    type=["mp3", "mp4", "mpeg", "mpga", "m4a", "wav", "webm"]
)


# =========================
# PROSES AUDIO
# =========================

if uploaded_file:

    st.audio(uploaded_file)
    st.write(f"**File:** {uploaded_file.name}")

    if st.button("🚀 Proses Audio"):

        try:
            with st.spinner("Memuat model transkripsi (Whisper)..."):
                model = load_whisper_model()

            suffix = os.path.splitext(uploaded_file.name)[1]
            with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as temp_file:
                temp_file.write(uploaded_file.getbuffer())
                temp_file_path = temp_file.name

            status_text = st.empty()
            progress_bar = st.progress(0)

            status_text.write("Menganalisis file audio...")
            audio = AudioSegment.from_file(temp_file_path)

            total_duration_ms = len(audio)
            chunk_length_ms = 30 * 1000
            num_chunks = max(1, -(-total_duration_ms // chunk_length_ms))

            transcript_parts = []
            start_time = time.time()

            for i in range(num_chunks):
                chunk_start = i * chunk_length_ms
                chunk_end = min((i + 1) * chunk_length_ms, total_duration_ms)
                chunk_audio = audio[chunk_start:chunk_end]

                with tempfile.NamedTemporaryFile(delete=False, suffix=".wav") as chunk_file:
                    chunk_audio.export(chunk_file.name, format="wav")
                    chunk_path = chunk_file.name

                result = model.transcribe(chunk_path, language="id")
                transcript_parts.append(result["text"])

                os.remove(chunk_path)

                progress = (i + 1) / num_chunks
                elapsed = time.time() - start_time
                estimated_total = elapsed / progress if progress > 0 else 0
                remaining = max(0, estimated_total - elapsed)

                progress_bar.progress(progress)
                status_text.write(
                    f"Mentranskripsi... bagian {i + 1}/{num_chunks} selesai "
                    f"(estimasi sisa waktu: {int(remaining)} detik)"
                )

            transcript = " ".join(transcript_parts).strip()
            os.remove(temp_file_path)

            progress_bar.progress(1.0)
            status_text.write("Transkripsi selesai!")

            with st.spinner("Membuat kesimpulan dan bahan belajar... (biasanya 10-30 detik)"):

                gemini_model = genai.GenerativeModel("gemini-3.6-flash")

                prompt = f"""
Kamu adalah asisten belajar untuk mahasiswa magang
yang sedang mempelajari materi teknikal.

Analisis transkrip meeting berikut.

Buat output dengan format:

1. Kesimpulan meeting
2. Poin-poin penting
3. Istilah teknikal dan penjelasan sederhananya
4. Hal yang perlu dipelajari lagi
5. Pertanyaan yang bisa ditanyakan ke mentor

Gunakan bahasa Indonesia yang natural,
jelas, dan mudah dipahami mahasiswa.

Jangan mengarang informasi yang tidak ada
di dalam transkrip.

TRANSKRIP:
{transcript}
"""

                response = gemini_model.generate_content(prompt)
                result_text = response.text

            st.session_state.transcript = transcript
            st.session_state.result_text = result_text

        except Exception as e:
            st.error(f"Terjadi error: {e}")


# =========================
# TAMPILKAN HASIL (dari session_state, tetap tampil walau di-rerun)
# =========================

if st.session_state.transcript and st.session_state.result_text:

    st.success("Transkripsi berhasil!")

    st.subheader("📝 Transkrip")
    st.text_area(
        "Hasil transkripsi",
        st.session_state.transcript,
        height=300
    )

    st.subheader("📚 Bahan Belajar")
    st.markdown(st.session_state.result_text)

    st.subheader("📥 Download Hasil")

    col1, col2 = st.columns(2)

    with col1:
        docx_buffer = generate_docx(st.session_state.transcript, st.session_state.result_text)
        st.download_button(
            "⬇️ Download sebagai Word (.docx)",
            data=docx_buffer,
            file_name="hasil_meeting.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

    with col2:
        pdf_buffer = generate_pdf(st.session_state.transcript, st.session_state.result_text)
        st.download_button(
            "⬇️ Download sebagai PDF",
            data=pdf_buffer,
            file_name="hasil_meeting.pdf",
            mime="application/pdf"
        )